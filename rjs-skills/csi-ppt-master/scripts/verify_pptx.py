#!/usr/bin/env python3
"""Independent structural check of a generated .pptx (does not trust the generator).

Checks per slide:
  - counts: text boxes, other shapes, tables, charts, pictures
  - text rendered as an image: a picture covering a large share of the slide = FAIL
  - fonts actually used by every run (latin / cs / ea typeface), incl. table cells
  - Thai text integrity: zero-width spaces, runs that start with a combining mark
    (a split tone mark / floating vowel), and the cs (complex-script) font on Thai runs
  - speaker notes present and non-empty
  - first line of the title-like text box (largest font on the slide) for action-title review

Usage:
  python3 verify_pptx.py deck.pptx [--expect-font Sarabun] [--json out.json]
Exit code: 0 = pass, 1 = fail (any FAIL item), 2 = cannot open file.
"""
from __future__ import annotations

import argparse
import json
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
THAI_RANGE = (0x0E00, 0x0E7F)
# Thai combining marks: above/below vowels and tone marks (must never start a run)
THAI_COMBINING = {0x0E31, *range(0x0E34, 0x0E3B), *range(0x0E47, 0x0E4F)}
ZERO_WIDTH = {"​", "‌", "‍", "⁠", "﻿"}
PICTURE_AREA_FAIL = 0.25  # a picture over 25% of the slide is treated as possible text-as-image


def has_thai(text: str) -> bool:
    return any(THAI_RANGE[0] <= ord(ch) <= THAI_RANGE[1] for ch in text)


def theme_fonts(prs) -> dict:
    """Map theme references (+mj-lt, +mn-cs, ...) to concrete typefaces of the first master's theme."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree

    out: dict[str, str] = {}
    try:
        theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)
    except Exception:  # noqa: BLE001 — a deck without a readable theme just keeps raw names
        return out
    for prefix, tag in (("+mj", "majorFont"), ("+mn", "minorFont")):
        font = root.find(f".//{A_NS}{tag}")
        if font is None:
            continue
        for slot, short in (("latin", "lt"), ("ea", "ea"), ("cs", "cs")):
            el = font.find(f"{A_NS}{slot}")
            if el is not None and el.get("typeface"):
                out[f"{prefix}-{short}"] = el.get("typeface")
        thai = font.find(f"{A_NS}font[@script='Thai']")
        if thai is not None and thai.get("typeface"):
            out[f"{prefix}-thai"] = thai.get("typeface")
    return out


THEME: dict[str, str] = {}


def resolve(name: str | None) -> str | None:
    if name and name.startswith("+"):
        return THEME.get(name, name)
    return name


def run_fonts(run) -> dict:
    """Return typefaces on the run (latin / ea / cs), theme references resolved; None when absent."""
    rpr = run._r.find(f"{A_NS}rPr")
    out = {"latin": None, "ea": None, "cs": None}
    if rpr is None:
        return out
    for tag in out:
        el = rpr.find(f"{A_NS}{tag}")
        if el is not None:
            out[tag] = resolve(el.get("typeface"))
    return out


def iter_shapes(shapes):
    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shp.shapes)
        else:
            yield shp


def text_frames_of(shp):
    if shp.has_text_frame:
        yield shp.text_frame
    if getattr(shp, "has_table", False) and shp.has_table:
        for row in shp.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def check_runs(frames, slide_no, findings, fonts_used, expect_font):
    max_size, title_text = 0, ""
    for tf in frames:
        for para in tf.paragraphs:
            for run in para.runs:
                text = run.text
                if not text:
                    continue
                f = run_fonts(run)
                for v in f.values():
                    if v:
                        fonts_used.add(v)
                if any(z in text for z in ZERO_WIDTH):
                    findings.append(("FAIL", slide_no, f"zero-width character inside text: {text[:30]!r}"))
                if ord(text[0]) in THAI_COMBINING:
                    findings.append(("FAIL", slide_no, f"run starts with a Thai combining mark (split cluster): {text[:30]!r}"))
                if has_thai(text):
                    cs = f["cs"] or f["latin"]
                    if expect_font and cs and cs != expect_font:
                        findings.append(("FAIL", slide_no, f"Thai run uses font {cs!r}, expected {expect_font!r}: {text[:30]!r}"))
                size = run.font.size.pt if run.font.size else 0
                if size > max_size:
                    max_size, title_text = size, para.text if hasattr(para, "text") else text
    return max_size, title_text


def verify(path: str, expect_font: str | None) -> dict:
    prs = Presentation(path)
    THEME.clear()
    THEME.update(theme_fonts(prs))
    slide_area = prs.slide_width * prs.slide_height
    report = {"file": path, "slides": [], "findings": []}
    findings = report["findings"]
    for idx, slide in enumerate(prs.slides, start=1):
        counts = {"text_boxes": 0, "shapes": 0, "tables": 0, "charts": 0, "pictures": 0}
        fonts_used: set[str] = set()
        frames = []
        big_pictures = 0
        for shp in iter_shapes(slide.shapes):
            if getattr(shp, "has_table", False) and shp.has_table:
                counts["tables"] += 1
                frames.extend(text_frames_of(shp))
            elif getattr(shp, "has_chart", False) and shp.has_chart:
                counts["charts"] += 1
            elif shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
                if shp.width and shp.height and (shp.width * shp.height) / slide_area > PICTURE_AREA_FAIL:
                    big_pictures += 1
            elif shp.has_text_frame and shp.text_frame.text.strip():
                counts["text_boxes"] += 1
                frames.append(shp.text_frame)
            else:
                counts["shapes"] += 1
        _, title = check_runs(frames, idx, findings, fonts_used, expect_font)
        if counts["text_boxes"] == 0 and counts["tables"] == 0:
            findings.append(("FAIL", idx, "no editable text on slide (text may be an image)"))
        if big_pictures:
            findings.append(("FAIL", idx, f"{big_pictures} picture(s) cover >{int(PICTURE_AREA_FAIL*100)}% of slide — check for text rendered as image"))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not notes:
            findings.append(("FAIL", idx, "speaker notes missing"))
        report["slides"].append({
            "slide": idx, **counts,
            "fonts": sorted(fonts_used),
            "title_candidate": title.strip()[:80],
            "notes_chars": len(notes),
        })
    report["pass"] = not any(level == "FAIL" for level, *_ in findings)
    return report


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--expect-font", help="font every Thai run must use (cs/latin typeface)")
    ap.add_argument("--json", help="write the full report as JSON to this path")
    args = ap.parse_args()
    try:
        rep = verify(args.pptx, args.expect_font)
    except Exception as exc:  # noqa: BLE001 — report any open/parse error plainly
        print(f"ERROR: cannot read {args.pptx}: {exc}", file=sys.stderr)
        return 2
    print("slide | text | shapes | tables | charts | pictures | notes | fonts | title")
    for s in rep["slides"]:
        print(f"{s['slide']:>5} | {s['text_boxes']:>4} | {s['shapes']:>6} | {s['tables']:>6} | {s['charts']:>6} | "
              f"{s['pictures']:>8} | {s['notes_chars']:>5} | {','.join(s['fonts']) or '-'} | {s['title_candidate']}")
    for level, slide, msg in rep["findings"]:
        print(f"[{level}] slide {slide}: {msg}")
    print("RESULT:", "PASS" if rep["pass"] else "FAIL")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as fh:
            json.dump(rep, fh, ensure_ascii=False, indent=2)
    return 0 if rep["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
